"""Analisa as planilhas oficiais de CEBAS (MS, MDS, MEC) publicadas no Mapa das OSCs
e cruza com os CNPJs extraidos do DOU pelo parser.

Uso:
    python cruzar_planilhas.py
Le:  ../mapa_osc/respostas/downloads/{1434-cebassaude,7684-cebassuas,8420-cebaseducacao}.xlsx
     cebas_S01*.csv (saida do parser_cebas.py)
Gera: cruzamento_planilhas.json
"""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl

AQUI = Path(__file__).parent
PLANILHAS = AQUI.parent / "mapa_osc" / "respostas" / "downloads"


def so_digitos(v) -> str:
    s = re.sub(r"\D", "", str(v or ""))
    return s.zfill(14) if 11 < len(s) <= 14 else s


def como_data(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, str):
        m = re.match(r"\s*(\d{2})/(\d{2})/(\d{4})", v)
        if m:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


def linhas(ws, linha_cabecalho: int) -> list[dict]:
    todas = list(ws.iter_rows(values_only=True))
    cab = [str(c).strip() if c is not None else f"col{i}" for i, c in enumerate(todas[linha_cabecalho])]
    return [dict(zip(cab, r)) for r in todas[linha_cabecalho + 1 :] if any(x is not None for x in r)]


def resumo_datas(registros: list[dict], colunas: list[str]) -> dict:
    out = {}
    for c in colunas:
        ds = [d for d in (como_data(r.get(c)) for r in registros) if d]
        if ds:
            out[c] = {"min": min(ds).isoformat(), "max": max(ds).isoformat(), "preenchidos": len(ds)}
    return out


def carregar() -> dict:
    res = {}

    wb = openpyxl.load_workbook(PLANILHAS / "1434-cebassaude.xlsx", read_only=True, data_only=True)
    ms = linhas(wb["TB_ENTIDADE_CEBAS"], 0)
    res["SAUDE"] = {
        "arquivo": "1434-cebassaude.xlsx",
        "aba": "TB_ENTIDADE_CEBAS",
        "registros": len(ms),
        "cnpjs": {so_digitos(r["NU_CNPJ"]) for r in ms},
        "colunas": list(ms[0].keys()),
        "situacoes": Counter(r["DS_SITUACAO_ATUAL"] for r in ms).most_common(),
        "st_cebas": Counter(r["ST_CEBAS"] for r in ms).most_common(),
        "datas": resumo_datas(ms, ["DT_INICIO_VIGENCIA", "DT_FIM_VIGENCIA"]),
        "por_cnpj": {so_digitos(r["NU_CNPJ"]): r for r in ms},
    }

    wb = openpyxl.load_workbook(PLANILHAS / "7684-cebassuas.xlsx", read_only=True, data_only=True)
    proc = linhas(wb["PRINCIPAL"], 4)
    sit = linhas(wb["SITUAÇÃO CNPJ CEBAS"], 4)
    res["ASSISTENCIA"] = {
        "arquivo": "7684-cebassuas.xlsx",
        "aba": "SITUAÇÃO CNPJ CEBAS (+ aba PRINCIPAL com processos)",
        "registros": len(sit),
        "registros_processos": len(proc),
        "cnpjs": {so_digitos(r["CNPJ"]) for r in sit},
        "cnpjs_processos": {so_digitos(r["CNPJ"]) for r in proc},
        "colunas": list(sit[0].keys()),
        "colunas_processos": list(proc[0].keys()),
        "situacoes": Counter(r["STATUS DE CERTIFICAÇÃO"] for r in sit).most_common(),
        "fases_processo": Counter(r["FASE_PROCESSO"] for r in proc).most_common(15),
        "datas": resumo_datas(sit, ["DT_INICIO_CERTIFICACAO_ATUAL", "DT_FIM_CERTIFICACAO_ATUAL"])
        | resumo_datas(proc, ["DT_PROTOCOLO", "DT_PUBLICACAO_PORTARIA_SNAS_DOU"]),
        "por_cnpj": {so_digitos(r["CNPJ"]): r for r in sit},
    }

    wb = openpyxl.load_workbook(PLANILHAS / "8420-cebaseducacao.xlsx", read_only=True, data_only=True)
    abas = wb.sheetnames
    mec = linhas(wb[abas[0]], 0)
    res["EDUCACAO"] = {
        "arquivo": "8420-cebaseducacao.xlsx",
        "aba": abas[0],
        "abas": abas,
        "registros": len(mec),
        "cnpjs": {so_digitos(r["CNPJ"]) for r in mec},
        "colunas": [k.strip() for k in mec[0].keys()],
        "situacoes": Counter(str(r.get("col3")) for r in mec).most_common(),
        "datas": resumo_datas(mec, [k for k in mec[0] if "DATA" in k or "VALIDADE" in k]),
        "por_cnpj": {so_digitos(r["CNPJ"]): r for r in mec},
    }
    return res


SISCEBAS_XLS = AQUI / "downloads" / "bases_abertas" / "siscebas_saude_20260930_ListaEntidadeSituacaoAtual.xls"


def analisar_siscebas_saude(dou: list[dict]) -> dict:
    """Lista publica do SisCEBAS Saude (gerada na hora, uma linha por requerimento)."""
    import xlrd

    # O XLS gerado pelo SisCEBAS tem o container OLE levemente corrompido; o conteudo le normal.
    wb = xlrd.open_workbook(str(SISCEBAS_XLS), ignore_workbook_corruption=True)
    sh = wb.sheet_by_index(0)
    cab = [str(c).strip() for c in sh.row_values(0)]
    regs = [dict(zip(cab, sh.row_values(i))) for i in range(1, sh.nrows)]
    for r in regs:
        r["_cnpj"] = so_digitos(r["CNPJ REQUERENTE"])
    cnpjs = {r["_cnpj"] for r in regs}
    com_cebas = {r["_cnpj"] for r in regs if r["CEBAS"] == "SIM"}
    pubs = [d for d in (como_data(r["DATA DA PUBLICAÇÃO"]) for r in regs) if d]

    dou_ms = {r["cnpj"]: r for r in dou if r["area"] == "SAUDE"}
    por_tipo, achou = Counter(), Counter()
    for c, r in dou_ms.items():
        k = f"{r['tipo_ato']}/{r['deferido'] or '-'}"
        por_tipo[k] += 1
        achou[k] += c in cnpjs

    # Atos do DOU em 2026 que a lista ja reflete (mesma data de publicacao)
    pub_por_cnpj: dict[str, set] = {}
    for r in regs:
        pub_por_cnpj.setdefault(r["_cnpj"], set()).add(como_data(r["DATA DA PUBLICAÇÃO"]))
    refletidos = sum(
        1
        for c, r in dou_ms.items()
        if r["data_publicacao"] and date.fromisoformat(r["data_publicacao"]) in pub_por_cnpj.get(c, set())
    )
    exemplos = {
        c: [{k: r[k] for k in cab} for r in regs if r["_cnpj"] == c]
        for c in ("50798453000183", "03163888000171")
    }
    return {
        "arquivo": SISCEBAS_XLS.name,
        "colunas": cab,
        "linhas_requerimentos": len(regs),
        "cnpjs_unicos": len(cnpjs),
        "cnpjs_com_cebas_sim": len(com_cebas),
        "situacoes_top": Counter(r["SITUAÇÃO ATUAL"] for r in regs).most_common(20),
        "cebas": Counter(r["CEBAS"] for r in regs).most_common(),
        "data_atualizacao": Counter(r["DATA ATUALIZAÇÃO"] for r in regs).most_common(3),
        "publicacao_mais_recente": max(pubs).isoformat() if pubs else None,
        "dou_ms_cnpjs": len(dou_ms),
        "dou_ms_na_lista": sum(1 for c in dou_ms if c in cnpjs),
        "dou_ms_por_tipo_na_lista": {k: f"{achou[k]}/{n}" for k, n in sorted(por_tipo.items())},
        "dou_ms_ato_refletido_mesma_data_publicacao": refletidos,
        "exemplos_t3": exemplos,
    }


def main() -> None:
    base = carregar()
    dou = []
    for f in sorted(AQUI.glob("cebas_S01*.csv")):
        dou += list(csv.DictReader(f.open(encoding="utf-8-sig"), delimiter=";"))
    dou = [r for r in dou if r["cnpj"] and not r["detalhe"].startswith("sem decisao")]

    todas = set().union(*(b["cnpjs"] for b in base.values()))
    cobertura = {}
    for area in ("SAUDE", "ASSISTENCIA", "EDUCACAO"):
        regs = [r for r in dou if r["area"] == area]
        cnpjs = {r["cnpj"] for r in regs}
        mesma = cnpjs & base[area]["cnpjs"]
        por_tipo = Counter()
        por_tipo_achou = Counter()
        for r in {r["cnpj"]: r for r in regs}.values():
            k = f"{r['tipo_ato']}/{r['deferido'] or '-'}"
            por_tipo[k] += 1
            por_tipo_achou[k] += r["cnpj"] in base[area]["cnpjs"]
        item = {
            "cnpjs_dou": len(cnpjs),
            "na_planilha_do_mesmo_ministerio": len(mesma),
            "pct": round(100 * len(mesma) / len(cnpjs), 1) if cnpjs else None,
            "em_qualquer_planilha": len(cnpjs & todas),
            "por_tipo_no_dou": {k: f"{por_tipo_achou[k]}/{n}" for k, n in sorted(por_tipo.items())},
        }
        if area == "ASSISTENCIA":
            item["na_aba_processos"] = len(cnpjs & base[area]["cnpjs_processos"])
        cobertura[area] = item

    # Exemplos T3
    exemplos = {}
    for cnpj in ("50798453000183", "03163888000171"):
        r = base["SAUDE"]["por_cnpj"].get(cnpj)
        exemplos[cnpj] = (
            {k: (v.isoformat() if isinstance(v, datetime) else v) for k, v in r.items()} if r else None
        )

    siscebas = analisar_siscebas_saude(dou)

    saida = {
        "siscebas_saude_lista_atual": siscebas,
        "planilhas": {
            a: {k: v for k, v in b.items() if k not in ("cnpjs", "cnpjs_processos", "por_cnpj")}
            | {"cnpjs_unicos": len(b["cnpjs"])}
            for a, b in base.items()
        },
        "cobertura_dou_jun_ago_2026": cobertura,
        "exemplos_t3_na_planilha_ms": exemplos,
    }
    (AQUI / "cruzamento_planilhas.json").write_text(
        json.dumps(saida, indent=2, ensure_ascii=False, default=str), encoding="utf-8"
    )
    print(json.dumps(saida, indent=2, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
